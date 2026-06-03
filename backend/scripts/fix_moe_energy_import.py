#!/usr/bin/env python
"""
Fix the MoE Energy mis-import: projects from the Ministry of Energy sheet landed
in Digital Transformation (as empty $0 shells) plus instruction/section banner
rows got imported as projects.

This script:
  1. Connects to whatever DATABASE_URL is in your environment (point it at PROD).
  2. Parses the MoE Energy .xlsx with the FIXED importer parser.
  3. Identifies the bad batch in Digital:
       - energy shells  : Digital projects whose name matches a row in the sheet
       - junk rows      : Digital projects whose name is an instruction/section banner
       - ag mislabel    : "Food systems resilience accelerator" (stays, re-mapped to Agriculture)
  4. DRY-RUN by default: prints exactly what it would delete / re-import / re-map.
     Pass --apply to execute (delete bad batch + dependents, re-import into the
     Energy TWG, re-map the Ag row) in a single transaction.

USAGE (run from backend/, with the PROD url in your own shell — keeps the secret
off the transcript):

    DATABASE_URL='postgresql://...proxy.rlwy.net:PORT/railway' \
        uv run python scripts/fix_moe_energy_import.py            # dry run
    DATABASE_URL='...' uv run python scripts/fix_moe_energy_import.py --apply

Optional: --sheet /path/to/MoE.xlsx  (defaults to the committed fixture).
"""
import argparse
import asyncio
import os
import sys
import uuid

import openpyxl
from sqlalchemy import text
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker

# Import the FIXED parser + pillar helpers from the app.
from app.api.routes.pipeline import parse_pipeline_workbook, _JUNK_NAME_MARKERS
from app.models.models import TWGPillar

DEFAULT_SHEET = os.path.join(
    os.path.dirname(__file__), "..", "tests", "fixtures", "import_sheets",
    "moe_energy_ministry_sheet.xlsx",
)
AG_MISLABEL_NAME = "Food systems resilience accelerator"

# Child tables that reference projects.id — clear these before deleting projects.
DEPENDENT_TABLES = [
    "project_status_history",
    "project_scores_detail",
    "project_buyer_matches",
    "project_investor_matches",
    "project_dfi_matches",
    "project_geospatial_data",
    "impact_log_entries",
]


def _is_junk(name: str) -> bool:
    nl = name.strip().lower()
    return any(m in nl for m in _JUNK_NAME_MARKERS)


async def main(apply: bool, sheet_path: str):
    db_url = os.environ.get("DATABASE_URL")
    if not db_url:
        sys.exit("ERROR: set DATABASE_URL (point it at the PROD database).")
    async_url = db_url.replace("postgresql://", "postgresql+asyncpg://")

    # Parse the sheet with the fixed parser using Energy as the fallback pillar.
    wb = openpyxl.load_workbook(sheet_path, data_only=True)
    parsed, skipped, errors = parse_pipeline_workbook(
        wb, fallback_pillar=TWGPillar.energy_infrastructure.value
    )
    sheet_names = {p["name"].strip() for p in parsed}
    print(f"Parsed sheet: {len(parsed)} projects, {skipped} skipped, {len(errors)} errors")
    if errors:
        print("  parse errors:", errors[:5])

    engine = create_async_engine(async_url)
    Session = async_sessionmaker(engine, expire_on_commit=False)
    async with Session() as db:
        host = (await db.execute(text("SELECT inet_server_addr()::text, current_database()"))).first()
        print(f"\nConnected to: host={host[0]} db={host[1]}")

        # Current Digital projects.
        digital = (await db.execute(text(
            "SELECT id::text, name FROM projects WHERE pillar = :p"
        ), {"p": TWGPillar.digital_economy_transformation.value})).all()
        print(f"Digital projects in DB: {len(digital)}")

        to_delete, junk, ag_remap, other_digital = [], [], [], []
        for pid, name in digital:
            nm = (name or "").strip()
            if nm == AG_MISLABEL_NAME:
                ag_remap.append((pid, nm))
            elif _is_junk(nm):
                junk.append((pid, nm))
                to_delete.append(pid)
            elif nm in sheet_names:
                to_delete.append(pid)
            else:
                other_digital.append((pid, nm))

        energy_shells = [pid for pid in to_delete if pid not in {j[0] for j in junk}]

        # Energy TWG (re-import target).
        twg = (await db.execute(text(
            "SELECT id::text, name FROM twgs WHERE pillar = :p ORDER BY id LIMIT 1"
        ), {"p": TWGPillar.energy_infrastructure.value})).first()

        print("\n================ PLAN ================")
        print(f"DELETE {len(to_delete)} bad rows:")
        print(f"   - {len(energy_shells)} energy shells (match sheet names)")
        print(f"   - {len(junk)} junk banner rows: {[j[1][:40] for j in junk]}")
        print(f"RE-IMPORT {len(parsed)} energy projects into TWG: "
              f"{twg[1] if twg else '!! NO energy_infrastructure TWG FOUND !!'} ({twg[0] if twg else '-'})")
        print(f"RE-MAP {len(ag_remap)} Ag row(s) -> agriculture_food_systems: {[a[1] for a in ag_remap]}")
        if other_digital:
            print(f"LEAVE UNTOUCHED {len(other_digital)} other Digital project(s): {[o[1][:40] for o in other_digital]}")
        print("======================================")

        if not apply:
            print("\nDRY RUN — nothing changed. Re-run with --apply to execute.")
            await engine.dispose()
            return

        if twg is None:
            sys.exit("ABORT: no TWG with pillar energy_infrastructure to import into.")

        # ---- APPLY (single transaction) ----
        async with db.begin():
            if to_delete:
                for tbl in DEPENDENT_TABLES:
                    try:
                        await db.execute(
                            text(f"DELETE FROM {tbl} WHERE project_id = ANY(:ids)"),
                            {"ids": to_delete},
                        )
                    except Exception as e:  # table may not exist in every env
                        print(f"  (skip {tbl}: {str(e)[:60]})")
                await db.execute(
                    text("DELETE FROM projects WHERE id = ANY(:ids)"), {"ids": to_delete}
                )

            for pid, _ in ag_remap:
                await db.execute(
                    text("UPDATE projects SET pillar = :p WHERE id = :id"),
                    {"p": TWGPillar.agriculture_food_systems.value, "id": pid},
                )

            cols = [
                "id", "twg_id", "name", "description", "investment_size", "currency",
                "status", "pillar", "lead_country", "subsector", "project_sponsor",
                "is_cross_border", "key_contact_name", "key_contact_email",
                "technical_studies", "permits_licences", "land_status",
                "financing_structure", "investment_stage_label", "revenue_model",
                "macroeconomic_roi", "climate_impact", "esg_compliance",
                "ghg_avoided_target", "jobs_construction", "jobs_om",
                "electricity_connections", "digital_connections",
                "smallholder_farmers_reached", "submitted_by",
            ]
            placeholders = ", ".join(f":{c}" for c in cols)
            insert_sql = text(
                f"INSERT INTO projects ({', '.join(cols)}) VALUES ({placeholders})"
            )
            for d in parsed:
                params = {c: None for c in cols}
                params.update(d)
                params["id"] = uuid.uuid4()
                params["twg_id"] = uuid.UUID(twg[0])
                # status is an enum object from the parser -> use its value
                st = params.get("status")
                params["status"] = st.value if hasattr(st, "value") else st
                await db.execute(insert_sql, params)

        print(f"\nAPPLIED: deleted {len(to_delete)}, re-imported {len(parsed)} into "
              f"{twg[1]}, re-mapped {len(ag_remap)} Ag row(s).")
        print("NOTE: re-imported projects come in at AfCEN score 0 / Draft — run a "
              "rescore pass to score them.")
    await engine.dispose()


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="execute (default is dry-run)")
    ap.add_argument("--sheet", default=DEFAULT_SHEET, help="path to the MoE .xlsx")
    args = ap.parse_args()
    asyncio.run(main(args.apply, args.sheet))
